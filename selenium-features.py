import tkinter as tk
from tkinter import scrolledtext, messagebox
from threading import Thread
import time
from datetime import datetime, timedelta
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, NoSuchElementException, StaleElementReferenceException
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.action_chains import ActionChains
import re
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font


class NewsScraperApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Course Scraper")
        self.center_window(800, 600)
        self.root.configure(bg='#1c1c1c')

        customFont = ('Consolas', 12)
        self.log = scrolledtext.ScrolledText(root, font=customFont, fg='#D8DEE9', bg='#2e2e2e', wrap=tk.WORD, padx=10, pady=10)
        self.log.grid(row=0, column=0, columnspan=3, padx=10, pady=10, sticky='nsew')

        root.grid_rowconfigure(0, weight=1)
        root.grid_columnconfigure(0, weight=1)

        buttonStyle = {'font': customFont, 'bg': '#5E81AC', 'fg': '#ECEFF4', 'padx': 10, 'pady': 5, 'relief': 'raised', 'bd': 3}
        self.start_button = tk.Button(root, text="Start Scraping", command=self.start_scraping, **buttonStyle)
        self.start_button.grid(row=1, column=0, sticky='ew', padx=10, pady=10)
        self.stop_button = tk.Button(root, text="Stop Scraping", command=self.stop_scraping, state='disabled', **buttonStyle)
        self.stop_button.grid(row=1, column=1, sticky='ew', padx=10, pady=10)

        self.running = False
        self.driver = None
        self.all_courses = []  # Track all scraped courses

        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)

    def center_window(self, width=800, height=600):
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        x = (screen_width // 2) - (width // 2)
        y = (screen_height // 2) - (height // 2)
        self.root.geometry(f'{width}x{height}+{x}+{y}')

    def log_message(self, message, tag=None):
        self.log.configure(state='normal')
        self.log.insert(tk.END, message + "\n", tag)
        self.log.configure(state='disabled')
        self.log.yview(tk.END)

    def initialize_driver(self):
        driver_path = r'C:\Mohamed\Personal-apps\Discount-courses-udemy-desktop-app\chromedriver_win64\chromedriver-win64\chromedriver.exe'
        service = Service(executable_path=driver_path)
        self.driver = webdriver.Chrome(service=service)
        self.wait = WebDriverWait(self.driver, 10)

    def is_course_recent(self, release_time_text):
        try:
            match = re.match(r"(\d+)\s+(minute|hour|day)s?\s+ago", release_time_text)
            if match:
                value, unit = int(match.group(1)), match.group(2)
                if unit == 'minute':
                    return True
                elif unit == 'hour':
                    return value < 24
                elif unit == 'day':
                    return value < 1
            elif "an hour ago" in release_time_text:
                return True
            return False
        except Exception as e:
            self.log_message(f"Error parsing release time: {release_time_text}. Exception: {str(e)}", 'error')
            return False

    def scrape_courses(self):
        if not self.driver:
            self.initialize_driver()

        try:
            self.log_message("Navigating to the website...", 'header')
            url = "https://www.real.discount"
            self.driver.get(url)
            time.sleep(2)

            self.log_message("Waiting for the '100% Off' checkbox to be visible...", 'header')
            free_checkbox = self.driver.find_element(By.CSS_SELECTOR, "label[for='free']")
            self.log_message("'100% Off' checkbox found, attempting to click...", 'header')
            try:
                isSelected = free_checkbox.is_selected()
                if not isSelected:
                    self.log_message("Checkbox was not clicked. Attempting ActionChains click.", 'header')
                    actions = ActionChains(self.driver)
                    actions.move_to_element(free_checkbox).click().perform()
                    self.log_message("Clicked the checkbox using ActionChains.", 'success')
            except Exception as e:
                self.log_message(f"Failed to click using ActionChains. Exception: {str(e)}", 'error')

            self.log_message("Waiting for the page to update after clicking...", 'header')
            time.sleep(2)

            self.log_message("Selecting the 'IT & Software' category...", 'header')
            category_dropdown = Select(self.driver.find_element(By.ID, 'category'))
            category_dropdown.select_by_visible_text('IT & Software')
            self.log_message("'IT & Software' category selected.", 'success')

            self.log_message("Waiting for the page to update after selecting category...", 'header')
            time.sleep(2)

            while self.running:
                self.log_message("Extracting course information...", 'header')
                courses = self.wait.until(EC.visibility_of_all_elements_located((By.TAG_NAME, 'li')))
                new_courses = []

                for course in courses:
                    try:
                        title_element = course.find_element(By.TAG_NAME, 'h3')
                        title = title_element.text if title_element else 'No Title'
                        link_element = course.find_element(By.TAG_NAME, 'a')
                        link = link_element.get_attribute('href') if link_element else 'No Link'
                        category = course.find_element(By.TAG_NAME, 'h5').text if course.find_elements(By.TAG_NAME, 'h5') else 'No Category'
                        release_time = course.find_element(By.CSS_SELECTOR, 'div.col-auto.ml-0.pl-0.pr-0.mr-0 > div').text if course.find_elements(By.CSS_SELECTOR, 'div.col-auto.ml-0.pl-0.pr-0.mr-0 > div') else 'No Release Time'

                        if title == 'No Title' or not link.startswith("https://www.real.discount") or '/out-ad/' in link:
                            continue

                        course_info = {
                            "Title": title,
                            "Category": category,
                            "Link": link,
                            "Release Time": release_time
                        }

                        if course_info not in self.all_courses:
                            self.all_courses.append(course_info)
                            new_courses.append(course_info)
                            self.log_message(f"Course Title: {title}", 'title')
                            self.log_message(f"Category: {category}", 'category')
                            self.log_message(f"Link: {link}", 'link')
                            self.log_message(f"Release Time: {release_time}", 'release_time')
                            self.log_message('-' * 40, 'separator')

                        if not self.is_course_recent(release_time):
                            self.log_message("Course older than 24 hours found. Stopping scraping.", 'header')
                            self.running = False
                            break

                    except NoSuchElementException:
                        continue
                    except Exception as course_e:
                        self.log_message(f"Error extracting data for a course: {course_e}", 'error')

                if not new_courses or not self.running:
                    self.log_message("No new courses found or stopping condition met. Stopping.", 'header')
                    break

                try:
                    load_more_button = self.wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, 'input.btn.btn-primary.mb-5')))
                    load_more_button.click()
                    self.log_message("Clicked the 'Load More' button.", 'success')
                    time.sleep(5)
                except TimeoutException:
                    self.log_message("No more 'Load More' button found. All courses loaded.", 'header')
                    break
                except Exception as e:
                    self.log_message(f"Failed to click 'Load More' button. Exception: {str(e)}", 'error')
                    break

        except TimeoutException as e:
            self.log_message(f"Timeout while waiting for an element. Exception: {str(e)}", 'error')
        except Exception as e:
            self.log_message(f"An error occurred: {str(e)}", 'error')
        finally:
            if not self.running:
                self.driver.quit()
                self.driver = None
                self.save_to_excel()

    def save_to_excel(self):
        if self.all_courses:
            df = pd.DataFrame(self.all_courses)
            try:
                # Create a filename with today's date
                today_date = datetime.now().strftime("%d%m%Y")
                file_name = f"free_courses_{today_date}.xlsx"
                df.to_excel(file_name, index=False)
                self.log_message(f"Scraped data saved to Excel file {file_name}.", 'success')

                # Load the workbook and select the active worksheet
                wb = load_workbook(file_name)
                ws = wb.active

                # Set column widths
                column_widths = {
                    "A": 50,  # Title
                    "B": 30,  # Category
                    "C": 50,  # Link
                    "D": 20   # Release Time
                }

                for col, width in column_widths.items():
                    col_letter = get_column_letter(ws[col + '1'].column)
                    ws.column_dimensions[col_letter].width = width

                # Apply bold font to header row
                header_font = Font(bold=True)
                for cell in ws[1]:
                    cell.font = header_font

                # Save the workbook with formatting
                wb.save(file_name)
                self.log_message(f"Formatted Excel file saved as {file_name}.", 'success')
            except Exception as e:
                self.log_message(f"Error saving to Excel file: {str(e)}", 'error')



    def start_scraping(self):
        self.running = True
        self.thread = Thread(target=self.scrape_courses)
        self.thread.start()
        self.start_button.config(state='disabled')
        self.stop_button.config(state='normal')
        self.log_message("Started scraping.", 'header')

    def stop_scraping(self):
        self.running = False
        if self.thread.is_alive():
            self.thread.join()
        self.start_button.config(state='normal')
        self.stop_button.config(state='disabled')
        self.log_message("Stopped scraping.", 'header')

    def on_closing(self):
        if messagebox.askokcancel("Quit", "Are you sure you want to quit?"):
            if self.running:
                self.stop_scraping()
            self.root.destroy()

def main():
    root = tk.Tk()
    app = NewsScraperApp(root)
    
    # Add custom tags for colored text
    app.log.tag_config('header', foreground='#88C0D0', font=('Consolas', 12, 'bold'))
    app.log.tag_config('title', foreground='#A3BE8C', font=('Consolas', 12, 'bold italic'))
    app.log.tag_config('category', foreground='#B48EAD', font=('Consolas', 12, 'italic'))
    app.log.tag_config('link', foreground='#EBCB8B', font=('Consolas', 12, 'underline'))
    app.log.tag_config('release_time', foreground='#D08770', font=('Consolas', 12))
    app.log.tag_config('separator', foreground='#4C566A', font=('Consolas', 10, 'normal'))
    app.log.tag_config('success', foreground='#A3BE8C')
    app.log.tag_config('error', foreground='#BF616A')
    
    root.mainloop()

if __name__ == "__main__":
    main()
